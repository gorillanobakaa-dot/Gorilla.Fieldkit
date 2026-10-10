"""fieldkit cv / fieldkit jobs: every rule is shown to fire on a bad input and to stay quiet on a good one.

The candidate is invented (Alex Example, @example.com, a 07700 900xxx number: Ofcom's range for drama). No real
person's data is in this file; the job boards are never called (a recorded response is passed in instead).
"""
import datetime as dt
import json
import zipfile

import pytest
import yaml

from fieldkit.career import advert, cv, pack, pipeline, scoring, sources, store, takehome
from fieldkit.career import profile as prof
from fieldkit.core import backup
from fieldkit.office.create import create

TODAY = dt.date(2026, 10, 10)

PROFILE = {
    "name": "Alex Example",
    "contact": {"phone": "07700 900123", "email": "alex@example.com", "location": "Birmingham"},
    "right_to_work": {"en": "Full right to work in the UK", "ro": "Drept deplin de muncă în Marea Britanie"},
    "years_experience": 7, "trade": "driver",
    "languages_spoken": ["English", "Romanian"],
    "languages": [{"language": {"en": "English", "ro": "Engleză"}, "level": {"en": "fluent", "ro": "fluent"}}],
    "licences": [{"category": "C+E", "detail": {"en": "Driver CPC, digital tachograph card",
                                                "ro": "CPC, card tahograf digital"}}],
    "salary_floor": 30000, "locations": ["Birmingham", "Coventry"],
    "exclude_title_patterns": ["trainee"],
    "experience": [
        {"title": {"en": "HGV Class 1 Driver", "ro": "Șofer HGV categoria C+E"}, "employer": "Example Haulage Ltd",
         "location": "Birmingham", "start": "2019-03", "end": "present",
         "bullets": [{"en": "Drove UK and EU routes for seven years without an accident.",
                      "ro": "Am condus pe rute în Marea Britanie și UE șapte ani fără niciun accident."},
                     {"text": {"en": "Loaded and checked 26-pallet trailers.", "ro": "Am încărcat remorci de 26 de paleți."},
                      "variants": ["driver"]},
                     {"text": {"en": "Trained two new warehouse staff.", "ro": "Am instruit doi colegi noi."},
                      "variants": ["warehouse"]}]},
        {"title": {"en": "Van Driver", "ro": "Șofer de dubă"}, "employer": "Sample Couriers", "start": "2017-01",
         "end": "2019-02", "bullets": [{"en": "Made 120 multi-drop deliveries a day.",
                                         "ro": "Am făcut 120 de livrări pe zi."}]},
    ],
    "cv_variants": {
        "driver": {"headline": {"en": "HGV Class 1 (C+E) Driver", "ro": "Șofer profesionist C+E"},
                   "profile": {"en": "Professional driver with seven years in England.",
                               "ro": "Șofer profesionist cu șapte ani de experiență în Anglia."},
                   "keywords": ["hgv", "class 1", "c+e", "driver", "multi-drop"],
                   "skills": [{"en": "Drivers' hours rules", "ro": "Reguli privind orele de condus"}]},
        "warehouse": {"headline": {"en": "Warehouse Operative", "ro": "Lucrător depozit"},
                      "profile": {"en": "Reliable warehouse operative.", "ro": "Lucrător de depozit."},
                      "keywords": ["warehouse", "picker", "forklift"],
                      "skills": [{"en": "Loading and stock checks", "ro": "Încărcare și verificare stoc"}]},
    },
    "evidence": [
        {"id": "safe", "first_person": "I have driven for seven years in England without an accident.",
         "source": "original CV"},
        {"id": "eu", "first_person": "I have driven HGV routes across France, Belgium and Germany.",
         "source": "candidate, 2026-10-10"},
        {"id": "multidrop", "first_person": "I have made up to 120 multi-drop deliveries a day.",
         "source": "original CV"},
    ],
    "letter": {"pitch": ["I hold a C+E licence, a Driver CPC and a digital tachograph card."]},
    "queries": [{"what": "HGV class 1 driver", "where": "Birmingham"}],
}


@pytest.fixture
def home(tmp_path, monkeypatch):
    """A data folder with the invented profile; prof.data_dir points at it."""
    d = tmp_path / "career"
    d.mkdir()
    (d / "profile.yaml").write_text(yaml.safe_dump(PROFILE, allow_unicode=True), encoding="utf-8")
    monkeypatch.setattr(prof, "data_dir", lambda local=None: d)
    return d


def write_docx(path, lines):
    create({"type": "docx", "blocks": [{"paragraph": l} for l in lines]}, path)
    return path


BAD_CV = ["Alex Example", "alex@example.com", "Date of birth: 01/01/1985", "Profile",
          "Hard working team player, responsible for various tasks etc.",
          "Experience", "Van Driver, Sample Couriers 2015 - 2017", "HGV Driver, Example Haulage 2019 - present",
          "Skills", "Driving the the lorry"]
GOOD_CV = ["Alex Example", "Birmingham | 07700 900123 | alex@example.com", "Profile",
           "Professional HGV driver with 7 years in England and EU routes. Clean licence, accident-free.",
           "Experience", "HGV Class 1 Driver, Example Haulage, Mar 2019 - Present",
           "On-time deliveries across 3 depots, 98% delivery windows kept; drivers' hours and tachograph rules.",
           "Van Driver, Sample Couriers, Jan 2017 - Feb 2019", "Licences and certificates",
           "Licence categories: B, C, C+E", "Driver CPC (DQC card), digital tachograph card, ADR, forklift (FLT)",
           "Skills", "Load security, route planning", "Languages", "English (fluent), Romanian (native)"]


# -- cv check ------------------------------------------------------------------------------------------------------------
def test_check_bad_cv_findings_in_order(tmp_path):
    r = cv.check(write_docx(tmp_path / "bad.docx", BAD_CV))
    ids = [f["id"] for f in r["findings"]]
    for must in ("no-phone", "date-of-birth", "not-newest-first", "no-licence-categories", "vague", "repeated-word",
                 "no-languages"):
        assert must in ids, must
    assert r["verdict"] == "needs work" and not r["ok"]
    sev = [f["severity"] for f in r["findings"]]
    assert sev == sorted(sev, key=lambda s: cv.SEVERITY[s])              # most important first
    assert r["trades"] == ["driver"]
    q = {x["id"] for x in r["questions"]}
    assert {"licence", "cpc", "tachograph", "adr", "forklift", "abroad", "other", "where", "pay"} <= q


def test_check_good_cv_is_quiet_on_what_it_has(tmp_path):
    r = cv.check(write_docx(tmp_path / "good.docx", GOOD_CV))
    ids = {f["id"] for f in r["findings"]}
    assert not ids & {"no-phone", "no-email", "date-of-birth", "not-newest-first", "no-licence-categories",
                      "no-cpc", "no-safety", "no-years", "vague", "no-experience"}
    assert r["licences"] == ["B", "C", "C+E"]
    lic = next(q for q in r["questions"] if q["id"] == "licence")
    assert lic["cv_says"] == "B, C, C+E"                                  # the question shows what the CV says


def test_another_trade(tmp_path):
    cvf = write_docx(tmp_path / "sec.docx", ["Alex Example", "07700 900123 alex@example.com", "Experience",
                                             "Security Officer, Example Guarding 2020 - present",
                                             "CCTV monitoring and patrols as a security officer"])
    r = cv.check(cvf)
    assert r["trades"] == ["security"] and "no-sia" in {f["id"] for f in r["findings"]}
    assert "no-licence-categories" not in {f["id"] for f in r["findings"]}   # a driver's rule, not a guard's


def test_photo_is_found(tmp_path):
    import docx
    from docx.shared import Cm
    from PIL import Image
    img = tmp_path / "p.png"
    Image.new("RGB", (10, 10)).save(img)
    d = docx.Document()
    d.add_paragraph("Alex Example")
    d.add_picture(str(img), width=Cm(2))
    d.save(tmp_path / "photo.docx")
    assert "photo" in {f["id"] for f in cv.check(tmp_path / "photo.docx")["findings"]}


# -- profile and render --------------------------------------------------------------------------------------------------
def test_profile_refuses(home):
    bad = dict(PROFILE, date_of_birth="1985-01-01")
    with pytest.raises(prof.BadProfile, match="date of birth"):
        prof.validate(bad)
    bad = dict(PROFILE, evidence=[{"id": "x", "first_person": "I am great."}])
    with pytest.raises(prof.BadProfile, match="source missing"):
        prof.validate(bad)
    exp = json.loads(json.dumps(PROFILE["experience"]))
    exp[0]["bullets"][1]["variants"] = ["pilot"]
    with pytest.raises(prof.BadProfile, match="variant 'pilot'"):
        prof.validate(dict(PROFILE, experience=exp))


def test_render_en_ro_ats_and_never_overwrites(home):
    from fieldkit.office.read import read
    r = cv.render(variant="driver", lang="en", today=TODAY)
    assert r["ok"], r["findings"]
    text = read(r["files"]["docx"])
    # the dates are apart from the job line in plain text (ATS): a real tab, not a positional one
    assert "Example Haulage Ltd, Birmingham\tMar 2019 – Present" in text
    assert "Trained two new warehouse staff" not in text                 # the other variant's bullet
    assert text.index("HGV Class 1 Driver") < text.index("Van Driver")    # newest first
    ro = cv.render(variant="driver", lang="ro", today=TODAY)
    pdf_text = read(ro["files"]["pdf"])
    assert "Experiență profesională".upper() in pdf_text and "Șofer" in pdf_text   # ș ț in the PDF
    again = cv.render(variant="driver", lang="en", today=TODAY)
    assert again["files"]["docx"].endswith("_2.docx") and r["files"]["docx"] != again["files"]["docx"]


def test_render_missing_translation_is_an_error(home):
    p = json.loads(json.dumps(PROFILE))
    p["cv_variants"]["driver"]["profile"] = {"en": "Only English."}
    (home / "profile.yaml").write_text(yaml.safe_dump(p, allow_unicode=True), encoding="utf-8")
    with pytest.raises(prof.BadProfile, match="no 'ro' text"):
        cv.render(variant="driver", lang="ro", today=TODAY)


# -- adverts and scoring -------------------------------------------------------------------------------------------------
AD = """HGV Class 1 Driver - Nights
Salary: £15.00 - £17.50 per hour
Closing date: 14 October 2026
Permanent, full-time. 2+ years experience required. Enhanced DBS.
"""


def test_advert_facts():
    f = advert.facts(AD, today=TODAY)
    assert (f["salary_min"], f["salary_max"], f["salary_period"]) == (15.0, 17.5, "hour")
    assert advert.pay_is_advertised(f) and f["closing_soon"] and f["years_required"] == 2
    assert "C+E" in f["licences_required"] and "DBS" in f["clearance"]
    est = advert.facts("Salary: £45,070 a year (estimated)\nFluent in Spanish required.", today=TODAY)
    assert est["salary_estimated"] and not advert.pay_is_advertised(est)
    assert est["languages_required"] == ["Spanish"]
    assert advert.facts("Closing date: 1 October 2026", today=TODAY)["closed"]


def job(**kw):
    j = {"key": "reed:1", "title": "HGV Class 1 Driver", "employer": "Example Logistics", "location": "Birmingham",
         "description": "C+E driver, multi-drop", "salary_min": 32000, "salary_max": 36000, "salary_predicted": 0,
         "advert_facts": None, "verified_salary_min": None}
    j.update(kw)
    return j


def test_scoring():
    p = prof.validate(PROFILE)
    good = scoring.score(job(), p)
    assert good["cv"] == "driver" and not good["excluded"] and good["score"] >= 70
    assert scoring.score(job(location="Leeds"), p)["score"] < good["score"]
    est = scoring.score(job(salary_predicted=1), p)
    assert any("ESTIMATE" in f for f in est["flags"]) and est["score"] < good["score"]
    low = scoring.score(job(salary_min=22000, salary_max=24000), p)
    assert any("below the floor" in f for f in low["flags"])
    pcv = scoring.score(job(title="PCV Bus Driver", description="bus driver"), p)
    assert pcv["excluded"] and "needs licence D" in pcv["excluded"]
    c_only = prof.validate(dict(PROFILE, licences=[{"category": "C"}]))
    assert "C+E" in scoring.score(job(), c_only)["excluded"]                 # Class 1 needs C+E
    assert scoring.score(job(title="Trainee HGV Driver"), p)["excluded"]
    lang = job(advert_facts=json.dumps(advert.facts("Fluent in German required.", today=TODAY)))
    assert "German" in scoring.score(lang, p)["excluded"]
    wh = scoring.score(job(title="Warehouse Picker", description="forklift warehouse picker"), p)
    assert wh["cv"] == "warehouse"


# -- sources: official APIs only -----------------------------------------------------------------------------------------
def test_sources_refuse_scraping_and_need_keys():
    for u in ("https://uk.indeed.com/jobs?q=driver", "https://www.linkedin.com/jobs", "https://example.com/api"):
        with pytest.raises(sources.NotAllowed):
            sources.check_url(u)
    with pytest.raises(sources.MissingKey, match="reed.co.uk/developers"):
        sources.reed({"what": "driver"}, {})
    assert "salary_min=30000" in sources.adzuna_url("i", "k", "driver", salary_min=30000)


def fake_opener(req):
    if "reed.co.uk" in req.full_url:
        assert req.headers["Authorization"].startswith("Basic ")
        return {"results": [{"jobId": 11, "jobTitle": "HGV Class 1 Driver", "employerName": "Example Logistics",
                             "locationName": "Birmingham", "minimumSalary": 33000, "maximumSalary": 36000,
                             "jobUrl": "https://www.reed.co.uk/jobs/11", "jobDescription": "C+E multi-drop driver"}]}
    return {"results": [{"id": "22", "title": "Class 1 Driver", "company": {"display_name": "Sample Freight"},
                         "location": {"display_name": "Coventry"}, "salary_min": 45070, "salary_max": 45070,
                         "salary_is_predicted": "1", "redirect_url": "https://www.adzuna.co.uk/jobs/22",
                         "description": "HGV class 1 driver"}]}


KEYS = {"reed_key": "r", "adzuna_app_id": "i", "adzuna_app_key": "k"}


# -- the whole cycle -----------------------------------------------------------------------------------------------------
def test_run_ingest_letter_status(home, monkeypatch):
    monkeypatch.setattr(sources, "keys", lambda: KEYS)
    p, t = prof.load(), store.Tracker(home / "jobs.sqlite")
    r = pipeline.run(p, t, top=8, opener=fake_opener, where=home)
    assert r["exit_code"] == 0 and r["new_jobs"] == 2 and len(r["shortlist"]) == 2
    adz = next(j for j in r["shortlist"] if j["key"] == "adzuna:22")
    assert adz["salary_status"] == "estimate" and any("ESTIMATE" in f for f in adz["flags"])
    assert (home / "packs" / "reed_11" / "cover_letter.md").is_file()
    assert json.loads((home / "reports" / "latest.json").read_text())["schema"] == "fieldkit.career.report/1"

    # the letter carries only approved sentences; anything added is caught
    j = t.get("reed:11")
    letter = (home / "packs" / "reed_11" / "cover_letter.md").read_text(encoding="utf-8")
    assert pack.check_letter(p, j, letter) == []
    assert pack.check_letter(p, j, letter + "\nI have managed a fleet of 200 lorries.\n") == \
        ["I have managed a fleet of 200 lorries."]

    # the estimate is not verified by an advert that does not state pay; a stated one is
    out = pipeline.ingest_advert(p, t, "adzuna:22", "Class 1 Driver\nCompetitive salary\n", where=home, today=TODAY)
    assert out["verdict"].startswith("pay NOT verified") and t.get("adzuna:22")["verified_at"] is None
    out = pipeline.ingest_advert(p, t, "adzuna:22", AD, where=home, today=TODAY)
    assert out["verdict"].startswith("pay VERIFIED") and t.get("adzuna:22")["verified_salary_max"] == 17.5 * 2080

    # a second run keeps statuses and adds nothing; a submitted job leaves the shortlist
    t.set_status("reed:11", "submitted_by_user", "candidate applied")
    r2 = pipeline.run(p, t, top=8, opener=fake_opener, where=home)
    assert r2["new_jobs"] == 0 and [j["key"] for j in r2["shortlist"]] == ["adzuna:22"]
    assert t.get("reed:11")["status"] == "submitted_by_user"
    hist = json.loads(t.get("reed:11")["history"])
    assert [h["status"] for h in hist][-1] == "submitted_by_user"


def test_run_without_keys_says_what_to_do(home, monkeypatch):
    monkeypatch.setattr(sources, "keys", lambda: {"reed_key": None, "adzuna_app_id": None, "adzuna_app_key": None})
    r = pipeline.run(prof.load(), store.Tracker(home / "jobs.sqlite"), where=home)
    assert r["exit_code"] == 4 and any("developer.adzuna.com" in h for h in r["human_actions"])


def test_one_failing_source_does_not_stop_the_other(home, monkeypatch):
    monkeypatch.setattr(sources, "keys", lambda: KEYS)

    def half(req):
        if "adzuna" in req.full_url:
            raise OSError("HTTP 500")
        return fake_opener(req)
    r = pipeline.run(prof.load(), store.Tracker(home / "jobs.sqlite"), opener=half, where=home)
    assert r["exit_code"] == 5 and r["new_jobs"] == 1


def test_verify_needs_a_source(home):
    p, t = prof.load(), store.Tracker(home / "jobs.sqlite")
    t.upsert(job(key="manual:1", source="manual", url="https://example.com/j/1"))
    with pytest.raises(ValueError, match="where the figures come from"):
        pipeline.verify(p, t, "manual:1", 30000, 34000)


def test_takehome():
    r = takehome.take_home(30000)
    assert (r["income_tax"], r["national_insurance"], r["take_home_year"]) == (3486.0, 1394.4, 25119.6)
    assert takehome.take_home(110000)["income_tax"] == 33432.0                # the allowance tapers above £100k
    assert r["tax_year"] == "2025/26"


# -- name, privacy, backup -----------------------------------------------------------------------------------------------
def test_init_asks_for_the_name_and_keeps_it_private(tmp_path):
    r = store.init(tmp_path / "c", name=None, lang="ro")
    assert r["status"] == "ask" and "Pe ce nume" in r["question"]
    assert not (tmp_path / "c").exists()                                  # nothing made before the answer
    r = store.init(tmp_path / "c", name="Alex Example", root=tmp_path)
    assert yaml.safe_load((tmp_path / "c" / "profile.yaml").read_text())["name"] == "Alex Example"
    terms = json.loads((tmp_path / "fieldkit.local.json").read_text())["privacy"]["terms"]
    assert "Alex Example" in terms and "Example" in terms
    store.init(tmp_path / "c", name="Alex Example", root=tmp_path)        # twice: the same list
    assert json.loads((tmp_path / "fieldkit.local.json").read_text())["privacy"]["terms"] == terms


def test_import_keeps_the_original_untouched(tmp_path):
    src = write_docx(tmp_path / "cv.docx", GOOD_CV)
    before = store.sha256(src)
    a = store.import_original(src, tmp_path / "c", today=TODAY)
    b = store.import_original(src, tmp_path / "c", today=TODAY)
    assert a["new"] and not b["new"] and a["copy"] == b["copy"] and store.sha256(src) == before


def test_backup(tmp_path):
    root = tmp_path / "Fieldkit"
    (root / "fieldkit").mkdir(parents=True)
    (root / "fieldkit" / "a.py").write_text("x = 1\n")
    (root / "local").mkdir()
    (root / "local" / "profile.yaml").write_text("name: x\n")
    (root / "toolbox").mkdir()
    (root / "toolbox" / "big.bin").write_bytes(b"0" * 1000)
    r = backup.backup(root=root, local={}, home=tmp_path)
    assert r["status"] == "ask"                                            # the folder is never guessed
    (tmp_path / "Backups").mkdir()
    assert str(tmp_path / "Backups") in backup.backup(root=root, local={}, home=tmp_path)["candidates"]
    r = backup.backup(to=tmp_path / "Backups", root=root, local={}, now=dt.datetime(2026, 10, 10, 18, 30))
    assert r["ok"] and r["archive"].endswith("harness_backup_2026-10-10_1830.zip") and r["files"] == 2
    names = zipfile.ZipFile(r["archive"]).namelist()
    assert "Fieldkit/local/profile.yaml" in names and not any("toolbox" in n for n in names)
    r2 = backup.backup(to=tmp_path / "Backups", root=root, local={}, now=dt.datetime(2026, 10, 10, 18, 30))
    assert r2["archive"].endswith("_2.zip")                                # never overwritten
    assert backup.backup(to=root / "inside", root=root, local={})["status"] == "refused"
    with zipfile.ZipFile(r["archive"], "a") as z:
        z.writestr("extra.txt", "x")
    assert not backup.verify(r["archive"], 2, r["bytes"])["ok"]           # a wrong archive is caught


def test_a_passing_mention_is_not_a_second_trade():
    """2026-10-10: a driver's CV listing 'forklift (FLT)' once was also checked as a warehouse CV."""
    assert cv.detect_trades("\n".join(GOOD_CV)) == ["driver"]
    both = "\n".join(["HGV driver", "Class 1 driver", "Warehouse picker", "Forklift (FLT) operator in a warehouse"])
    assert cv.detect_trades(both) == ["driver", "warehouse"]
