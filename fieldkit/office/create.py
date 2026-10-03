"""Create Word, Excel, PowerPoint and PDF files from a JSON/YAML spec.

The agent writes data, Fieldkit writes the file. The same spec always gives
the same document, and every file written is re-opened and checked before
this returns (check.py), so a broken file is an error here, not a surprise
when someone opens it.

Spec shapes (one "type" per file):

  docx: {type: docx, title?, blocks: [
            {heading: "text", level: 1} | {paragraph: "text"} |
            {bullets: ["a", "b"]} | {numbered: [...]} |
            {table: [["h1","h2"],["a","b"]]} | {page_break: true}]}
  xlsx: {type: xlsx, sheets: [{name, rows: [[...]], widths?: {A: 20},
            header?: true, freeze?: "A2"}]}
        A cell string starting with "=" is written as a formula.
  pptx: {type: pptx, slides: [{title, bullets?: [...], notes?}]}
  pdf:  {type: pdf, title?, blocks: [{heading}|{paragraph}|{bullets}]}

Metadata: author/last-modified-by are written EMPTY unless the spec sets
"author" - the file should not carry the machine owner's name by accident.
"""
import os
import tempfile
from pathlib import Path

from . import check as _check


def _docx(spec, out):
    import docx
    from docx.enum.text import WD_BREAK
    d = docx.Document()
    if spec.get("title"):
        d.add_heading(spec["title"], level=0)
    for b in spec.get("blocks", []):
        if "heading" in b:
            d.add_heading(b["heading"], level=int(b.get("level", 1)))
        elif "paragraph" in b:
            d.add_paragraph(b["paragraph"])
        elif "bullets" in b:
            for item in b["bullets"]:
                d.add_paragraph(item, style="List Bullet")
        elif "numbered" in b:
            for item in b["numbered"]:
                d.add_paragraph(item, style="List Number")
        elif "table" in b:
            rows = b["table"]
            t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    t.cell(i, j).text = "" if v is None else str(v)
        elif b.get("page_break"):
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            raise ValueError(f"unknown docx block: {b}")
    cp = d.core_properties
    cp.author = spec.get("author", "")
    cp.last_modified_by = spec.get("author", "")
    cp.title = spec.get("title", "")
    d.save(out)


def _xlsx(spec, out):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in spec.get("sheets", []):
        ws = wb.create_sheet(sh.get("name", "Sheet1")[:31])
        for row in sh.get("rows", []):
            ws.append(row)
        if sh.get("header", True) and ws.max_row >= 1:
            for c in ws[1]:
                c.font = Font(bold=True)
        for col, w in (sh.get("widths") or {}).items():
            ws.column_dimensions[col].width = w
        if sh.get("freeze"):
            ws.freeze_panes = sh["freeze"]
    wb.properties.creator = spec.get("author", "")
    wb.properties.lastModifiedBy = spec.get("author", "")
    wb.save(out)


def _pptx(spec, out):
    from pptx import Presentation
    prs = Presentation()
    for s in spec.get("slides", []):
        layout = prs.slide_layouts[1 if s.get("bullets") else 5]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = s.get("title", "")
        if s.get("bullets"):
            body = slide.placeholders[1].text_frame
            body.text = s["bullets"][0]
            for item in s["bullets"][1:]:
                body.add_paragraph().text = item
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    prs.core_properties.author = spec.get("author", "")
    prs.core_properties.last_modified_by = spec.get("author", "")
    prs.save(out)


def _pdf(spec, out):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape
    st = getSampleStyleSheet()
    flow = []
    if spec.get("title"):
        flow += [Paragraph(escape(spec["title"]), st["Title"]), Spacer(1, 8)]
    for b in spec.get("blocks", []):
        if "heading" in b:
            flow.append(Paragraph(escape(b["heading"]), st[f"Heading{min(int(b.get('level', 1)), 3)}"]))
        elif "paragraph" in b:
            flow.append(Paragraph(escape(b["paragraph"]), st["BodyText"]))
        elif "bullets" in b:
            flow.append(ListFlowable([ListItem(Paragraph(escape(i), st["BodyText"])) for i in b["bullets"]],
                                     bulletType="bullet"))
        else:
            raise ValueError(f"unknown pdf block: {b}")
    doc = SimpleDocTemplate(str(out), pagesize=A4, title=spec.get("title", ""), author=spec.get("author", ""),
                            creator="fieldkit", invariant=1)   # invariant=1: byte-identical output
    doc.build(flow)


WRITERS = {"docx": _docx, "xlsx": _xlsx, "pptx": _pptx, "pdf": _pdf}


def create(spec, out, force=False):
    """Write `out` from `spec`, then check it. Returns the check report.

    An existing `out` is never overwritten unless force=True. The file is
    written under a temporary name in the same folder and checked there; only a
    file that passes its own check is moved to `out`, so a failed check leaves
    no bad file behind (and an existing file untouched).
    """
    out = Path(out)
    ext = out.suffix.lstrip(".").lower()
    kind = spec.get("type") or ext
    if kind not in WRITERS:
        raise ValueError(f"type must be one of {sorted(WRITERS)}, got {kind!r}")
    if ext != kind:
        raise ValueError(f"{out.name}: spec type is {kind!r} but the file name ends .{ext}; "
                         f"name the output .{kind} or change the spec's type")
    if out.exists() and not force:
        raise FileExistsError(f"{out} already exists; refusing to overwrite it (use --force to replace it)")
    out.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".fieldkit-create-", suffix=out.suffix, dir=out.parent)
    os.close(fd)
    tmp = Path(tmp)
    try:
        WRITERS[kind](spec, tmp)
        report = _check.check(tmp)
        if report["problems"]:
            raise RuntimeError(f"{out.name} failed its own check and was not written: {report['problems']}")
        if out.exists() and not force:          # appeared while we were writing
            raise FileExistsError(f"{out} already exists; refusing to overwrite it (use --force to replace it)")
        os.replace(tmp, out)
    finally:
        if tmp.exists():
            tmp.unlink()
    report["file"] = str(out)
    return report
