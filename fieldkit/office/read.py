"""Read Word, Excel, PowerPoint and PDF files into Markdown-ish text, offline.

Uses the open, locally installed libraries (python-docx, openpyxl,
python-pptx; pdftotext from Poppler or pypdfium2 for PDF). If Microsoft's
markitdown or IBM's Docling is installed it can be chosen with
engine="markitdown" / "docling"; the built-in engine needs neither and is
what the tests pin down. Nothing is uploaded anywhere.
"""
import subprocess
from pathlib import Path

from ..core.host import find_tool

SUPPORTED = (".docx", ".xlsx", ".xlsm", ".pptx", ".pdf")


def _docx(path):
    import docx
    d = docx.Document(path)
    out = []
    body = d.element.body
    # Walk paragraphs and tables in document order.
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            if style.startswith("Heading"):
                level = "".join(c for c in style if c.isdigit()) or "1"
                out.append("#" * min(int(level), 6) + " " + text)
            elif style == "Title":
                out.append("# " + text)
            elif "List" in style:
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            t = Table(child, d)
            rows = [[c.text.strip().replace("|", "\\|") for c in r.cells] for r in t.rows]
            if rows:
                out.append(_md_table(rows))
    return "\n\n".join(out)


def _md_table(rows):
    """One Markdown table as ONE block: rows on consecutive lines (a blank line between rows ends the table)."""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        rows = [["" if v is None else str(v).replace("|", "\\|") for v in row]
                for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            out.append("(empty)")
            continue
        out.append(_md_table(rows))
    wb.close()
    return "\n\n".join(out)


def _pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        out.append(f"## Slide {i}" + (f": {title}" if title else ""))
        for shape in slide.shapes:
            if shape == slide.shapes.title or not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t:
                    out.append("  " * para.level + "- " + t)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"> Notes: {notes}")
    return "\n".join(out)


def _pdf(path):
    tool = find_tool("pdftotext")
    if tool:
        r = subprocess.run([tool, "-layout", "-enc", "UTF-8", str(path), "-"], capture_output=True)
        if r.returncode == 0:
            return r.stdout.decode("utf-8", "replace")
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(path))
    try:
        return "\n\f".join(doc[i].get_textpage().get_text_range() for i in range(len(doc)))
    finally:
        doc.close()


def read(path, engine="builtin"):
    """Return the document's text (Markdown for Office files)."""
    path = Path(path)
    ext = path.suffix.lower()
    if engine == "markitdown":
        from markitdown import MarkItDown
        return MarkItDown().convert(str(path)).text_content
    if engine == "docling":
        from docling.document_converter import DocumentConverter
        return DocumentConverter().convert(str(path)).document.export_to_markdown()
    if ext == ".docx":
        return _docx(path)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext == ".pdf":
        return _pdf(path)
    raise ValueError(f"{path.name}: unsupported type {ext}; supported: {', '.join(SUPPORTED)}")
