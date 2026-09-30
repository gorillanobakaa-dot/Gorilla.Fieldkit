"""Check that an Office or PDF file is structurally sound, before anyone opens it.

Our own open implementation of the checks that matter most in practice
(Anthropic's document skills carry XSD validators, but under a licence that
forbids copying, so nothing of theirs is used here):

Office (.docx/.xlsx/.pptx are ZIP packages of XML parts):
  - it is a readable ZIP, with no duplicate entries and no path tricks
  - [Content_Types].xml and _rels/.rels exist
  - every XML part is well-formed
  - every internal relationship points at a part that exists
  - the main part (word/document.xml, xl/workbook.xml, ppt/presentation.xml) exists
  - the file opens in the matching Python library
  - xlsx: formulas are listed, and formulas with no cached value are flagged
    (Excel shows them, but other readers see blanks until recalculated)
PDF:
  - header, EOF marker, opens, page count, whether text can be extracted
"""
import posixpath
import zipfile
from pathlib import Path

from lxml import etree

MAIN_PART = {".docx": "word/document.xml", ".docm": "word/document.xml",
             ".xlsx": "xl/workbook.xml", ".xlsm": "xl/workbook.xml",
             ".pptx": "ppt/presentation.xml", ".pptm": "ppt/presentation.xml"}
_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)


def _rels_source(rels_name):
    """'word/_rels/document.xml.rels' -> 'word/document.xml'; '_rels/.rels' -> ''."""
    folder, base = posixpath.split(rels_name)
    parent = posixpath.dirname(folder)
    return posixpath.join(parent, base[:-5]) if base != ".rels" else ""


def check_office(path):
    path = Path(path)
    problems, info = [], {}
    try:
        z = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        return {"file": str(path), "problems": ["not a ZIP package (corrupt, or not really an Office file)"], "info": {}}
    with z:
        names = z.namelist()
        info["parts"] = len(names)
        if len(set(names)) != len(names):
            problems.append("duplicate ZIP entries")
        for n in names:
            if n.startswith("/") or ".." in n.split("/"):
                problems.append(f"unsafe part name: {n}")
        for req in ("[Content_Types].xml", "_rels/.rels"):
            if req not in names:
                problems.append(f"missing {req}")
        main = MAIN_PART.get(path.suffix.lower())
        if main and main not in names:
            problems.append(f"missing main part {main}")
        present = set(names)
        for n in names:
            if not (n.endswith(".xml") or n.endswith(".rels")):
                continue
            try:
                root = etree.fromstring(z.read(n), _PARSER)
            except etree.XMLSyntaxError as e:
                problems.append(f"malformed XML in {n}: {e}")
                continue
            if n.endswith(".rels"):
                src_dir = posixpath.dirname(_rels_source(n))
                for rel in root:
                    if rel.get("TargetMode") == "External":
                        continue
                    tgt = rel.get("Target", "")
                    full = tgt.lstrip("/") if tgt.startswith("/") else posixpath.normpath(posixpath.join(src_dir, tgt))
                    if full not in present:
                        problems.append(f"{n}: relationship {rel.get('Id')} points at missing part {full}")
    # Library round-trip.
    ext = path.suffix.lower()
    try:
        if ext in (".docx", ".docm"):
            import docx
            d = docx.Document(path)
            info["paragraphs"] = len(d.paragraphs)
            info["tables"] = len(d.tables)
        elif ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path)                 # formulas as written
            wbv = openpyxl.load_workbook(path, data_only=True)  # cached values
            formulas, uncached = 0, []
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and c.value.startswith("="):
                            formulas += 1
                            if wbv[ws.title][c.coordinate].value is None:
                                uncached.append(f"{ws.title}!{c.coordinate}")
            info.update(sheets=wb.sheetnames, formulas=formulas, formulas_without_cached_value=len(uncached))
            if uncached:
                info["uncached_examples"] = uncached[:10]
        elif ext in (".pptx", ".pptm"):
            from pptx import Presentation
            info["slides"] = len(Presentation(path).slides)
    except Exception as e:  # noqa: BLE001 - any failure to open is the finding
        problems.append(f"does not open in the Python library: {type(e).__name__}: {e}")
    return {"file": str(path), "problems": problems, "info": info}


def check_pdf(path):
    path = Path(path)
    problems, info = [], {}
    data = path.read_bytes()
    if not data.startswith(b"%PDF-"):
        problems.append("no %PDF- header")
    if b"%%EOF" not in data[-2048:]:
        problems.append("no %%EOF marker near the end (truncated?)")
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(path))
        info["pages"] = len(doc)
        chars = sum(len(doc[i].get_textpage().get_text_range()) for i in range(min(len(doc), 5)))
        info["text_chars_first_5_pages"] = chars
        info["has_text"] = chars > 0
        doc.close()
    except Exception as e:  # noqa: BLE001
        problems.append(f"does not open: {type(e).__name__}: {e}")
    return {"file": str(path), "problems": problems, "info": info}


def check(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return check_pdf(path)
    if ext in MAIN_PART:
        return check_office(path)
    raise ValueError(f"{Path(path).name}: cannot check {ext}")
