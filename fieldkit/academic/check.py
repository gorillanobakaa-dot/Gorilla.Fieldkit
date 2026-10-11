"""The whole pre-submission check of one piece of work, in order, with one verdict; delivery only when it passes.

    fieldkit academic check FOLDER DRAFT [--offline]

DRAFT is the markdown draft, or the Word or PowerPoint file written straight into a template. Every check reads what
was written, in the student's own settings (profile.json: work language, spelling, referencing style, whether long
dashes count against the work). Each gate is PASS, FAIL, WARN (look at it) or SKIP (not checked: never a pass).

     1  saved on this computer  for a profile that keeps the work off the cloud (support.refuse_cloud), a folder in
                                OneDrive, Dropbox or Google Drive FAILS and stops the run: nothing is written there,
                                not even the report
     2  study materials read    every file in the materials folder read into the library; an unreadable one is named
     3  template leftovers      [placeholders], grey guidance paragraphs, comments and tracked changes (they carry the
                                author's name)
     4  references              each citation has its entry, each entry is cited, no duplicates, the list in order;
                                a quotation without its author, year and page is a WARN
     5  referencing style       entries and citations in the shape the style asks for (WARN)
     6  long dashes             a FAIL where the country's rules count them against the work, else a WARN
     7  language                a sentence in another language FAILS; the other spelling variant WARNS
     8  structure               the sections the document type asks for (SKIP while the type is not set)
     9  word count              the assessed words within 10% of the target (SKIP without a target)
    10  copied text             12 or more words verbatim from a study source FAIL; 7 to 11, or long runs identical
                                to the student's own notes, WARN; no sources is a SKIP, never a pass
    11  identifiers online      every DOI, ISBN, PMID ... resolves, to the work the entry names (--offline: SKIP)
    12  anonymous marking       the student's name is not in the text, the header or the file name
    13  delivered               only when nothing FAILED: a markdown draft is built, a Word or PowerPoint file is
                                copied and finished (header, page numbers, language, empty author fields), into
                                "4 - HERE IS YOUR WORK". The student's own file is never changed

Born from the academic harness this generalises (run_pipeline.py, 2026-10-10): the verdict comes first and the one
thing to fix first under it, because the end of a long report is the worst place for the answer; the same list is
written as a file to tick off (drafts folder), and `fieldkit academic next` reads the last verdict.
"""
import datetime as _dt
import os
import re
import shutil

from . import profile as prof

PLACEHOLDER = re.compile(
    r"\[(?:Your [^\]\n]{1,40}|Title|Student ID|Module Code|Key point|TODO(?:[:\s][^\]\n]{0,60})?|"
    r"Insert [^\]\n]{1,40}|References in [^\]\n]{1,40})\]")

FIX = {
    "saved on this computer": "move the assignment off the cloud folder first: fieldkit academic local --apply",
    "template leftovers": "fill every [ ] and delete the grey guidance; in Word, Review tab: accept or reject every "
                          "tracked change and delete every comment",
    "references": 'fieldkit academic refs "FILE" names each citation without an entry, each entry never cited, '
                  "the duplicates and the order",
    "long dashes": "replace each long dash (— –) with a comma, a colon or brackets: fieldkit academic dashes \"FILE\"",
    "language": 'rewrite the sentences that are not in the work language: fieldkit academic language "FILE"',
    "structure": "add the sections the type asks for: fieldkit academic types TYPE",
    "word count": "bring the assessed words within 10% of the target",
    "copied text": "put each run of 12+ words in quotation marks with its citation and page, or write it in your "
                   "own words: fieldkit academic plagiarism \"FILE\" --sources FOLDER",
    "identifiers online": "an identifier in the reference list does not exist or leads to another work: check it "
                          "against the source itself",
    "anonymous marking": "take your name out of the text, the header and the file name; the Student ID identifies "
                         "the work",
}


def _gate(gates, name, status, detail=""):
    gates.append({"name": name, "status": status, "detail": detail})


def read_config(folder):
    """{type, words, module} from EXAM_CONFIG.md, then the assignment state."""
    info = {"type": None, "words": None, "module": None}
    path = os.path.join(folder, "EXAM_CONFIG.md")
    text = open(path, encoding="utf-8", errors="replace").read() if os.path.isfile(path) else ""
    m = re.search(r"\*\*Document type:\*\*\s*(\S[^\n]*)", text)
    if m and m.group(1).strip().upper() not in ("TBD", "NONE", "-"):
        info["type"] = m.group(1).strip()
    m = re.search(r"\*\*Word count target:\*\*\s*([\d,]+)", text)
    if m:
        info["words"] = int(m.group(1).replace(",", ""))
    m = re.search(r"\*\*Module:\*\*\s*(\S[^\n]*)", text)
    if m and m.group(1).strip().upper() != "TBD":
        info["module"] = m.group(1).strip()
    from .assignment import load_state
    st = load_state(folder) or {}
    for k in info:
        info[k] = info[k] or st.get(k)
    return info


def _read(draft):
    """(markdown text, speaker notes, template info or None) of the draft, whatever it is written in."""
    low = draft.lower()
    if low.endswith(".docx"):
        from .build_templates import docx_to_markdown
        text, info = docx_to_markdown(draft)
        return text, "", info
    from .reference_auditor import load_document
    text, notes = load_document(draft)
    return text, notes, None


def _header_text(draft):
    if not draft.lower().endswith(".docx"):
        return ""
    try:
        from docx import Document
        return " ".join(p.text for s in Document(draft).sections for p in s.header.paragraphs)
    except Exception:
        return ""


def anonymity(text, name, filename, header=""):
    """The places the student's name appears: the file name (any part of it), the text or the header (the full
    name, so a cited author who shares one name does not trip it)."""
    if not name or not name.strip():
        return []
    parts = [p for p in re.split(r"\s+", name.strip()) if len(p) > 2]
    leaks = []
    base = os.path.basename(filename).lower()
    if any(p.lower() in base for p in parts):
        leaks.append("the file name")
    if name.strip().lower() in text.lower():
        leaks.append("the text")
    if header and name.strip().lower() in header.lower():
        leaks.append("the page header")
    return leaks


def run(folder, draft, offline=False, today=None, deliver=True):
    folder = os.path.abspath(folder)
    p = prof.load() or {}
    gates = []
    cfg = read_config(folder)
    out = {"folder": folder, "draft": os.path.abspath(draft), "type": cfg["type"], "target_words": cfg["words"]}

    # 1. on this computer: before anything is read or written
    from .workspace import Workspace, cloud_service, support
    svc = cloud_service(folder)
    if svc and support(p).get("refuse_cloud"):
        _gate(gates, "saved on this computer", "FAIL", f"in {svc}; your profile keeps the work off the cloud")
        return _finish(out, gates, folder, None, write=False)
    _gate(gates, "saved on this computer", "PASS", f"in {svc}: it syncs (the profile allows it)" if svc else "")
    if not os.path.isdir(folder):
        raise FileNotFoundError(f"no assignment folder {folder}: fieldkit academic setup FOLDER first")
    if not os.path.isfile(draft):
        raise FileNotFoundError(f"no such file: {draft}")
    if not draft.lower().endswith((".md", ".txt", ".docx", ".pptx")):
        raise ValueError("the draft is a .md, .docx or .pptx file")
    ws = Workspace(folder, p).create()

    # 2. the study materials
    try:
        from .extract_library import build_library
        lib = build_library(folder, verbose=False)
        infos, failed = lib["infos"], [rel for rel, _ in lib["failed"]]
        if not infos:
            _gate(gates, "study materials read", "WARN", "no study files in the folder: nothing to compare against")
        elif failed:
            _gate(gates, "study materials read", "WARN", f"{len(failed)} of {len(infos)} unreadable: "
                                                         + ", ".join(failed[:3]))
        else:
            _gate(gates, "study materials read", "PASS", f"{len(infos)} file(s), "
                                                         f"{sum(i['words'] for i in infos):,} words")
    except Exception as e:                                    # a broken reader never stops the check
        _gate(gates, "study materials read", "WARN", f"could not read them: {e}"[:200])

    text, notes, tinfo = _read(draft)
    stem = re.sub(r"^TEMPLATE\s*-\s*", "", os.path.splitext(os.path.basename(draft))[0], flags=re.I).strip() \
        or "submission"
    os.makedirs(ws.drafts, exist_ok=True)
    checked = os.path.join(ws.drafts, f"{stem} (as checked).md")
    if not draft.lower().endswith((".md", ".txt")):
        with open(checked, "w", encoding="utf-8") as f:
            f.write(text)
    else:
        checked = draft
    out["checked"] = checked

    # 3. template leftovers
    header = _header_text(draft)
    left = PLACEHOLDER.findall(re.sub(r"<!--.*?-->", "", text, flags=re.S)) + PLACEHOLDER.findall(header)
    grey = (tinfo or {}).get("guidance_left", 0)
    marks = 0
    if draft.lower().endswith(".docx"):
        from .office_finalise import review_marks
        marks = review_marks(draft)[0]
    parts = [f"{len(left)} [to fill]: {', '.join(sorted(set(left))[:4])}" if left else "",
             f"{grey} grey guidance paragraph(s)" if grey else "",
             f"{marks} comment(s) or tracked change(s)" if marks else ""]
    parts = [x for x in parts if x]
    _gate(gates, "template leftovers", "FAIL" if parts else "PASS", "; ".join(parts) or "none")

    # 4. references and quotations
    from .reference_auditor import audit_text
    audit = audit_text(text, notes)
    out["references"] = {k: audit[k] for k in ("total_citations", "total_references", "orphans", "unused",
                                               "duplicates", "out_of_order")}
    if audit["issues"]:
        bits = [f"{len(audit['orphans'])} cited, not listed" if audit["orphans"] else "",
                f"{len(audit['unused'])} listed, never cited" if audit["unused"] else "",
                f"{len(audit['duplicates'])} duplicate(s)" if audit["duplicates"] else "",
                f"{len(audit['out_of_order'])} out of order" if audit["out_of_order"] else ""]
        _gate(gates, "references", "FAIL", "; ".join(b for b in bits if b))
    else:
        _gate(gates, "references", "PASS", f"{audit['total_citations']} citation(s), "
                                           f"{audit['total_references']} entr(ies)")
    if audit["quote_warnings"]:
        _gate(gates, "quotations", "WARN", f"{len(audit['quotes_unattributed'])} without author and year, "
                                           f"{len(audit['quotes_without_page'])} without a page: add them by hand")
    else:
        _gate(gates, "quotations", "PASS", f"{audit['quotes']} quotation(s), each with its source")

    # 5. referencing style
    style = p.get("style")
    if style:
        from . import styles
        try:
            st = styles.check(text, style)
            _gate(gates, "referencing style", "WARN" if st["warnings"] else "PASS",
                  f"{st['warnings']} not in the {st['label']} shape" if st["warnings"] else st["label"])
        except (KeyError, ValueError) as e:
            _gate(gates, "referencing style", "SKIP", str(e)[:120])
    else:
        _gate(gates, "referencing style", "SKIP", "no style in the profile: fieldkit academic init")

    # 6. long dashes
    from .dash_check import find_dashes
    faults, dash_warn = find_dashes(text)
    counts = p.get("avoid_long_dashes", True)
    if faults:
        _gate(gates, "long dashes", "FAIL" if counts else "WARN", f"{len(faults)} long dash(es)")
    elif dash_warn:
        _gate(gates, "long dashes", "WARN", f"{len(dash_warn)} hyphen(s) used as a dash")
    else:
        _gate(gates, "long dashes", "PASS", "none")

    # 7. language
    from . import language
    lang = p.get("work_language") or "en"
    spelling = p.get("spelling") if lang == "en" else None
    lr = language.check(text, lang, spelling)
    if lr["other"]:
        _gate(gates, "language", "FAIL", f"{len(lr['other'])} sentence(s) in another language "
                                         f"({', '.join(sorted({o['name'] for o in lr['other']}))})")
    elif lr["spelling"]:
        _gate(gates, "language", "WARN", f"{len(lr['spelling'])} word(s) in the other spelling: "
              + ", ".join(f"{s['word']} -> {s['expected']}" for s in lr["spelling"][:4]))
    else:
        _gate(gates, "language", "PASS", lang + (f" ({spelling})" if spelling else ""))

    # 8. structure
    from . import document_types as dt
    key = dt.resolve_type(cfg["type"]) if cfg["type"] else None
    if not cfg["type"]:
        _gate(gates, "structure", "SKIP", "the document type is not set: fieldkit academic brief FOLDER")
    elif not key:
        _gate(gates, "structure", "FAIL", f"unknown document type {cfg['type']!r}")
    else:
        sr = dt.check_structure(text, key)
        if sr["failures"]:
            _gate(gates, "structure", "FAIL", f"{len(sr['failures'])} required part(s) missing: "
                                              + "; ".join(sr["failures"][:2]))
        elif sr["warnings"]:
            _gate(gates, "structure", "WARN", f"{len(sr['warnings'])} thing(s) to check")
        else:
            _gate(gates, "structure", "PASS", dt.DOCUMENT_TYPES[key]["label"])

    # 9. word count
    from .word_counter import TOLERANCE, WORD_RE, extract_body
    words = len(WORD_RE.findall(extract_body(text)))
    out["words"] = words
    if cfg["words"]:
        lo, hi = int(cfg["words"] * (1 - TOLERANCE)), int(cfg["words"] * (1 + TOLERANCE))
        _gate(gates, "word count", "PASS" if lo <= words <= hi else "FAIL",
              f"{words:,} words" + ("" if lo <= words <= hi else f"; between {lo:,} and {hi:,} asked"))
    else:
        _gate(gates, "word count", "SKIP", f"{words:,} words; no target set")

    # 10. copied text
    has_sources = os.path.isdir(ws.materials) and any(
        f for _, _, fs in os.walk(ws.materials) for f in fs if not f.upper().startswith("README"))
    if not has_sources:
        _gate(gates, "copied text", "SKIP", "no study materials to compare against: this is NOT a pass")
    else:
        from .plagiarism_precheck import is_serious, run_precheck
        matches = run_precheck(checked, ws.materials)
        serious = [m for m in matches if is_serious(m)]
        notes_long = [m for m in matches if not m["quoted"] and m.get("role") == "notes" and m["length"] >= 12]
        moderate = [m for m in matches if not m["quoted"] and m.get("role", "source") == "source"
                    and 7 <= m["length"] < 12]
        if serious:
            _gate(gates, "copied text", "FAIL", f"{len(serious)} run(s) of 12+ words copied exactly from "
                                                + ", ".join(sorted({m['source'] for m in serious})[:3]))
        elif notes_long:
            _gate(gates, "copied text", "WARN", f"{len(notes_long)} long run(s) identical to your own notes: "
                                                "check they are not copied")
        elif moderate:
            _gate(gates, "copied text", "WARN", f"{len(moderate)} run(s) of 7 to 11 words: look at each")
        else:
            _gate(gates, "copied text", "PASS", "nothing copied")

    # 11. identifiers online
    if offline:
        _gate(gates, "identifiers online", "SKIP", "--offline: run again with the internet before submitting")
    else:
        from .external_reference_validator import internet_available, run_validation
        if not internet_available():
            _gate(gates, "identifiers online", "SKIP", "no internet: run again before submitting")
        else:
            import contextlib
            import io
            with contextlib.redirect_stdout(io.StringIO()):
                ok = run_validation(checked, offline=False, verbose=False)
            _gate(gates, "identifiers online", "PASS" if ok else "FAIL",
                  "" if ok else "an identifier does not exist or leads to another work")

    # 12. anonymous marking
    name = p.get("name")
    leaks = anonymity(text, name, draft, header)
    if not name:
        _gate(gates, "anonymous marking", "SKIP", "no name in the profile to look for")
    else:
        _gate(gates, "anonymous marking", "FAIL" if leaks else "PASS",
              ("the name is in " + ", ".join(leaks)) if leaks else "not in the text, header or file name")

    # 13. delivered
    delivered = None
    blocking = [g for g in gates if g["status"] == "FAIL"]
    if not deliver:
        _gate(gates, "delivered", "SKIP", "not asked to deliver")
    elif blocking:
        _gate(gates, "delivered", "SKIP", f"nothing goes into {os.path.basename(ws.output)} until every FAIL is "
                                          "fixed; your file is unchanged")
    else:
        delivered = _deliver(draft, ws, stem, key)
        _gate(gates, "delivered", "PASS", delivered)
    return _finish(out, gates, folder, ws, delivered=delivered, today=today)


def _free(path):
    base, ext = os.path.splitext(path)
    n = 2
    while os.path.exists(path):
        path = "%s (%d)%s" % (base, n, ext)
        n += 1
    return path


def _deliver(draft, ws, stem, key):
    from . import build as b
    os.makedirs(ws.output, exist_ok=True)
    if draft.lower().endswith((".md", ".txt")):
        kind = "docx"
        if key:
            from .document_types import builder_kind, output_format
            kind = "poster" if builder_kind(key) == "poster" else ("pptx" if output_format(key) == "pptx" else "docx")
        return b.build(draft, kind, ws.output)["file"]
    target = _free(os.path.join(ws.output, stem + os.path.splitext(draft)[1].lower()))
    shutil.copy2(draft, target)                              # the student's own file is never changed
    b.finish(target)
    return target


def _finish(out, gates, folder, ws, delivered=None, today=None, write=True):
    failed = [g for g in gates if g["status"] == "FAIL"]
    warned = [g for g in gates if g["status"] == "WARN"]
    skipped = [g for g in gates if g["status"] == "SKIP"]
    verdict = ("NOT READY" if failed else "READY TO SUBMIT" if not warned else "READY, AFTER YOU LOOK AT THE WARNINGS")
    first = failed[0] if failed else None
    out.update({"ok": not failed, "verdict": verdict, "gates": gates, "delivered": delivered,
                "counts": {"pass": len(gates) - len(failed) - len(warned) - len(skipped), "fail": len(failed),
                           "warn": len(warned), "skip": len(skipped)},
                "first_fix": {"gate": first["name"], "detail": first["detail"],
                              "fix": FIX.get(first["name"], first["detail"]).replace("FILE", out["draft"])}
                if first else None})
    if write and ws is not None:
        out["todo"] = _todo(ws, out, today)
        from .assignment import _save_state, load_state
        st = load_state(folder) or {}
        when = today if isinstance(today, _dt.datetime) else _dt.datetime.now()
        st["last_check"] = {"when": when.strftime("%Y-%m-%d %H:%M"), "verdict": verdict,
                            "failed": [g["name"] for g in failed], "todo": out["todo"]}
        _save_state(folder, st)
    out["next"] = (out["first_fix"]["fix"] if failed else
                   f"open {delivered} and read it once more, then upload it" if delivered else
                   f'fieldkit academic check "{folder}" "{out["draft"]}"')
    return out


def _todo(ws, out, today):
    path = os.path.join(ws.drafts, "TO FIX.md")
    lines = [f"# {out['verdict']}", "", f"Checked: {os.path.basename(out['draft'])}", ""]
    for g in out["gates"]:
        box = "x" if g["status"] == "PASS" else " "
        lines.append(f"- [{box}] {g['status']:<4} {g['name']}" + (f": {g['detail']}" if g["detail"] else ""))
        if g["status"] == "FAIL" and g["name"] in FIX:
            lines.append(f"      fix: {FIX[g['name']].replace('FILE', out['draft'])}")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    return path


def lines(r):
    out = [r["verdict"]]
    if r.get("first_fix"):
        out.append(f"START WITH THIS (1 of {r['counts']['fail']}): {r['first_fix']['gate']}: {r['first_fix']['detail']}")
        out.append(f"  {r['first_fix']['fix']}")
    out.append("")
    for g in r["gates"]:
        out.append(f"  {g['status']:<4}  {g['name']:<24} {g['detail']}")
    c = r["counts"]
    out.append(f"{c['pass']} pass, {c['fail']} fail, {c['warn']} warn, {c['skip']} skipped"
               + (f"; words: {r['words']:,}" if r.get("words") is not None else ""))
    if r.get("todo"):
        out.append(f"the same list to tick off: {r['todo']}")
    return out
