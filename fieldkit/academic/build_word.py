"""
DOCUMENT BUILDER - Turns the report markdown into a submission-ready Word file.

The layout comes from layout.py (country, referencing style, work language,
profile); with no profile it is the UK default the harness was built on:
  * Arial 12pt, 1.5 line spacing (APA 7: double), black text
  * 2.54cm margins all round, A4 (US: Letter)
  * Anonymous header: Student ID and module code only, never a name,
    labelled in the work language
  * Page numbers in the footer
  * Hanging-indent, alphabetically ordered reference list
  * A page break before the reference list

Markdown actually supported (the old version passed most of this through as
literal text, so bullets arrived as "- item" and *italics* kept their stars):
  * # / ## / ### headings
  * **bold**, *italic*, ***bold italic***, `code`
  * - / * / + bullet lists, and 1. numbered lists, including nested levels
  * > block quotations (indented, per Harvard long-quote convention)
  * | pipe | tables |
  * --- horizontal rules (rendered as a page break)
  * Blank lines separate paragraphs; they no longer each become an empty
    paragraph, which was doubling the line spacing of the whole document.

Usage:
    fieldkit academic build DRAFT.md
"""
import os
import re
import sys
import argparse

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BODY_FONT = "Arial"
BODY_SIZE = Pt(12)
HEADING_SIZES = {1: Pt(16), 2: Pt(14), 3: Pt(12)}

from .reference_auditor import REF_HEADING                 # every language's heading

# Set by layout.apply(); DEFAULT here is the UK layout, so the module works on its own.
LAYOUT = None


def _lay():
    if LAYOUT is None:
        from .layout import current
        return current({})
    return LAYOUT


def _spacing_rule():
    s = float(_lay()["line_spacing"])
    return {1.0: WD_LINE_SPACING.SINGLE, 1.5: WD_LINE_SPACING.ONE_POINT_FIVE,
            2.0: WD_LINE_SPACING.DOUBLE}.get(s, WD_LINE_SPACING.ONE_POINT_FIVE)


# --------------------------------------------------------------------------
# Document scaffolding
# --------------------------------------------------------------------------

def _set_base_style(doc):
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE
    style.font.color.rgb = RGBColor(0, 0, 0)
    # Ensure East Asian and complex-script runs also use Arial, or Word
    # substitutes a different face for any non-Latin character.
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), BODY_FONT)

    pf = style.paragraph_format
    pf.line_spacing_rule = _spacing_rule()
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)          # paragraph spacing, not blank paragraphs


def _set_margins(doc):
    lay = _lay()
    width, height = lay["page_cm"]
    margin = Cm(float(lay["margins_cm"]))
    for section in doc.sections:
        # python-docx starts from a US Letter template; the country decides
        # (A4 everywhere but the US), or the page reflows for the marker.
        section.page_width = Cm(width)
        section.page_height = Cm(height)
        section.top_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin
        section.right_margin = margin


def page_furniture(path):
    """What the pages carry: (header text, are they numbered?).

    A document written by hand in Word often has neither. Under anonymous
    marking the Student ID on the header is what identifies the submission,
    and most handbooks ask for page numbers.
    """
    from docx import Document as _Doc
    doc = _Doc(path)
    section = doc.sections[0]
    header = " ".join(p.text for p in section.header.paragraphs).strip()
    numbered = any("PAGE" in p._p.xml for p in section.footer.paragraphs) or \
        any("PAGE" in p._p.xml for p in section.header.paragraphs)
    return header, numbered


def add_page_furniture(path, student_id, module_code="", force=False,
                       a4=True):
    """Fix the page setup of a document that already exists.

    Adds the Student ID header and page numbers if they are missing, and
    puts the page on A4 if Word left it on US Letter. Anything already
    there is left alone unless force is set. Only the page setup, header
    and footer are touched; the body is not read or rewritten.
    """
    from docx import Document as _Doc
    doc = _Doc(path)
    header, numbered = page_furniture(path)
    added = []
    if force or not header:
        _set_header(doc, student_id, module_code or "[Module Code]")
        added.append("header")
    if force or not numbered:
        _add_page_number_footer(doc)
        added.append("page numbers")

    section = doc.sections[0]
    width, height = _lay()["page_cm"]
    if a4 and round(section.page_width.cm, 1) != round(width, 1):
        for s in doc.sections:
            s.page_width = Cm(width)
            s.page_height = Cm(height)
        added.append("%s page size" % _lay()["page"])
        # The page numbers in an existing contents list are now stale, so
        # ask Word to rebuild its fields when the document is opened.
        settings = doc.settings.element
        update = settings.find(qn("w:updateFields"))
        if update is None:
            update = OxmlElement("w:updateFields")
            settings.append(update)
        update.set(qn("w:val"), "true")

    if added:
        doc.save(path)
    return added


def _set_header(doc, student_id, module_code):
    """Anonymous marking: Student ID and module only, never a name."""
    lab = _lay()["labels"]
    hp = doc.sections[0].header.paragraphs[0]
    hp.text = "%s: %s | %s: %s" % (lab["student_id"], student_id, lab["module"], module_code)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.name = BODY_FONT
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(128, 128, 128)


def _add_page_number_footer(doc):
    """Insert a live PAGE field, so numbering updates as the text changes."""
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    run.font.name = BODY_FONT
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(128, 128, 128)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# --------------------------------------------------------------------------
# Inline markdown
# --------------------------------------------------------------------------

INLINE_RE = re.compile(
    r"(\*\*\*.+?\*\*\*"      # bold italic
    r"|\*\*.+?\*\*"          # bold
    r"|__.+?__"              # bold (underscore)
    r"|\*[^*\n]+?\*"         # italic
    r"|_[^_\n]+?_"           # italic (underscore)
    r"|`[^`\n]+?`)"          # inline code
)


def add_inline(paragraph, text, size=None, bold=False, italic=False):
    """Add text to a paragraph, honouring inline markdown emphasis."""
    size = size or BODY_SIZE
    for part in INLINE_RE.split(text):
        if not part:
            continue
        run_bold, run_italic, mono = bold, italic, False
        content = part

        if part.startswith("***") and part.endswith("***"):
            content, run_bold, run_italic = part[3:-3], True, True
        elif part.startswith("**") and part.endswith("**"):
            content, run_bold = part[2:-2], True
        elif part.startswith("__") and part.endswith("__"):
            content, run_bold = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            content, run_italic = part[1:-1], True
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            content, run_italic = part[1:-1], True
        elif part.startswith("`") and part.endswith("`"):
            content, mono = part[1:-1], True

        # Escaped markdown and stray link syntax.
        content = re.sub(r"\\([*_`#\[\]])", r"\1", content)
        content = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", content)

        run = paragraph.add_run(content)
        run.bold = run_bold
        run.italic = run_italic
        run.font.name = "Consolas" if mono else BODY_FONT
        run.font.size = size
    return paragraph


# --------------------------------------------------------------------------
# Block-level markdown
# --------------------------------------------------------------------------

def _is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def _is_table_divider(line):
    return bool(re.match(r"^\s*\|[\s:\-|]+\|\s*$", line))


def _parse_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _add_table(doc, rows):
    if not rows:
        return
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], cell_text, size=Pt(11), bold=True)

    for row in body:
        cells = table.add_row().cells
        for i, cell_text in enumerate(row[:len(header)]):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], cell_text, size=Pt(11))


# Markdown level -> Word style. "#" is the document title; "##" sections are
# Heading 1, so the contents page and the navigation pane list the sections.
HEADING_STYLES = {1: "Title", 2: "Heading 1", 3: "Heading 2",
                  4: "Heading 3", 5: "Heading 3", 6: "Heading 3"}
STYLE_SIZES = {"Title": Pt(16), "Heading 1": Pt(14), "Heading 2": Pt(12),
               "Heading 3": Pt(12)}


def _force_font(style_or_run_element, name=None):
    """Set a font on every script, and strip theme fonts that override it.

    Word's built-in heading styles point at theme fonts (asciiTheme etc.).
    A theme font silently wins over a plain font name, which is why headings
    would otherwise come out in Calibri Light whatever the style says.
    """
    name = name or BODY_FONT
    rpr = style_or_run_element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme",
                 "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _set_heading_styles(doc):
    """The layout on Word's real heading styles: the body font, black, bold."""
    for style_name, size in STYLE_SIZES.items():
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = BODY_FONT
        style.font.size = size
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        _force_font(style.element)
        pf = style.paragraph_format
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        pf.line_spacing_rule = _spacing_rule()
    try:
        title = doc.styles["Title"]
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # The default Title style carries a coloured bottom border.
        ppr = title.element.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)
    except KeyError:
        pass


def _heading(doc, text, level, centred=False):
    style_name = HEADING_STYLES.get(level, "Heading 3")
    try:
        p = doc.add_paragraph(style=style_name)
    except KeyError:
        p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if centred:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, text, size=STYLE_SIZES.get(style_name, BODY_SIZE), bold=True)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def _add_contents_page(doc):
    """A Word table of contents built from the Heading styles.

    Word fills it in when the document is opened (updateFields is set), or
    on right-click > Update Field.
    """
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(12)
    run = label.add_run(_lay()["labels"]["contents"])
    run.bold = True
    run.font.name = BODY_FONT
    run.font.size = Pt(14)

    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)

    placeholder = p.add_run(_lay()["labels"]["contents_hint"])
    placeholder.italic = True
    placeholder.font.name = BODY_FONT

    end_run = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)

    # Ask Word to refresh fields on opening, so the contents appear.
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    _page_break(doc)
    return p


def _bullet(doc, text, depth, ordered, number):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.75 + 0.75 * depth)
    pf.first_line_indent = Cm(-0.4)
    pf.space_after = Pt(3)
    marker = ("%d. " % number) if ordered else (u"\u2022  " if depth == 0 else u"\u2013  ")
    run = p.add_run(marker)
    run.font.name = BODY_FONT
    run.font.size = BODY_SIZE
    add_inline(p, text)
    return p


def _blockquote(doc, text):
    """Long quotations: an indented block, no quote marks (every style here)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(1.27)
    pf.right_indent = Cm(1.27)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    add_inline(p, text, size=Pt(11))
    return p


def _reference_entry(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(1.27)
    pf.first_line_indent = Cm(-1.27)
    pf.space_after = Pt(6)
    add_inline(p, text)
    return p


def _page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def _ensure_page_break(doc):
    """A page break, unless the document already ends with one.

    Several things ask for a fresh page - the title, the contents, the
    reference list - and two of them in a row would leave a blank page.
    """
    paragraphs = doc.paragraphs
    if paragraphs:
        last = paragraphs[-1]
        if not last.text.strip() and any(
                b.get(qn("w:type")) == "page"
                for b in last._p.findall(".//" + qn("w:br"))):
            return None
    return _page_break(doc)


def _split_entries(lines):
    """Reference-list lines -> entries, markdown kept; an indented or non-entry line continues the one above."""
    from .reference_auditor import NAME_WORD, YEAR, YEAR_OR_ND, _strip_markdown
    entries = []
    for raw in lines:
        plain = _strip_markdown(raw).strip()
        starts = (not re.match(r"^[ \t]{2,}", raw)
                  and re.match(r"^[\-\*\d.\s]*" + NAME_WORD, plain)
                  and re.search(r"\(\s*" + YEAR_OR_ND + r"\s*\)|\b" + YEAR + r"\b", plain, flags=re.I))
        text = re.sub(r"^\s*[-*+]\s+", "", raw.strip())
        if starts or not entries:
            entries.append(text)
        else:
            entries[-1] += " " + text
    return [e for e in entries if e.strip()]


def _sort_key(entry):
    """Alphabetical order, ignoring markdown, leading articles and accents (the reference check's own order)."""
    from .reference_auditor import sort_key
    return sort_key(re.sub(r"[*_`]", "", entry))


# --------------------------------------------------------------------------
# Builder
# --------------------------------------------------------------------------

def _real(value):
    """A property value, or '' for a template placeholder like [Student ID]."""
    return "" if not value or "[" in value else value


def build_word_document(md_path, output_path, student_id="[Student ID]",
                        module_code="[Module Code]", sort_references=True,
                        verbose=True, contents_page=False, title_page=True):
    with open(md_path, "r", encoding="utf-8", errors="replace") as f:
        raw = f.read()

    raw = re.sub(r"<!--.*?-->", "", raw, flags=re.DOTALL)
    lines = raw.splitlines()

    doc = Document()
    _set_base_style(doc)
    _set_margins(doc)
    _set_header(doc, student_id, module_code)
    _add_page_number_footer(doc)
    _set_heading_styles(doc)

    # A contents page goes where the draft asks for one ("## Contents" or
    # "## Table of Contents"), or straight after the title if requested.
    contents_heading = re.compile(r"^#{1,6}\s+(?:(?:table\s+of\s+)?contents|[íi]ndice|inhaltsverzeichnis"
                                  r"|cuprins)\s*$", re.IGNORECASE)
    draft_has_contents = any(contents_heading.match(l.strip()) for l in lines)
    contents_pending = [contents_page and not draft_has_contents]
    contents_made = [False]

    in_references = False
    reference_entries = []
    paragraph_buffer = []
    table_buffer = []
    ordered_counters = {}

    def flush_paragraph():
        """Emit the buffered wrapped lines as one paragraph."""
        if not paragraph_buffer:
            return
        if in_references:
            # One entry per line that starts like an entry (the reference
            # check's own rule), so entries need no blank line between them.
            reference_entries.extend(_split_entries(paragraph_buffer))
            paragraph_buffer.clear()
            return
        text = " ".join(paragraph_buffer).strip()
        paragraph_buffer.clear()
        if not text:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, text)

    def flush_table():
        if table_buffer:
            _add_table(doc, list(table_buffer))
            table_buffer.clear()

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()

        # Tables are buffered until the block ends.
        if _is_table_row(line):
            flush_paragraph()
            if not _is_table_divider(line):
                table_buffer.append(_parse_row(line))
            continue
        flush_table()

        if not stripped:
            flush_paragraph()
            ordered_counters.clear()
            continue

        # Horizontal rule -> page break.
        if re.match(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$", line):
            flush_paragraph()
            _page_break(doc)
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2).strip()
            ordered_counters.clear()

            if contents_heading.match(stripped):
                if not contents_made[0]:
                    _ensure_page_break(doc)
                    _add_contents_page(doc)
                    contents_made[0] = True
                continue

            if re.match(r"^#*\s*" + REF_HEADING + r"\s*:?$", stripped,
                        flags=re.IGNORECASE):
                in_references = True
                _ensure_page_break(doc)
                _heading(doc, text, min(level, 2))
                continue

            if level >= 2 and contents_pending[0]:
                _ensure_page_break(doc)
                _add_contents_page(doc)
                contents_pending[0] = False
                contents_made[0] = True

            _heading(doc, text, level, centred=(level == 1))
            # The title sits on a page of its own (a cover page); whatever
            # comes next starts a new page. title_page=False turns it off.
            if level == 1 and title_page:
                _ensure_page_break(doc)
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            _blockquote(doc, stripped.lstrip("> ").strip())
            continue

        bullet = re.match(r"^(\s*)([-*+])\s+(.*)$", line)
        if bullet:
            flush_paragraph()
            depth = len(bullet.group(1)) // 2
            _bullet(doc, bullet.group(3).strip(), depth, False, 0)
            continue

        numbered = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if numbered:
            flush_paragraph()
            depth = len(numbered.group(1)) // 2
            ordered_counters[depth] = ordered_counters.get(depth, 0) + 1
            _bullet(doc, numbered.group(3).strip(), depth, True,
                    ordered_counters[depth])
            continue

        # Ordinary prose: buffer it so wrapped lines join into one paragraph.
        paragraph_buffer.append(line if in_references else stripped)

    flush_paragraph()
    flush_table()

    if reference_entries:
        entries = sorted(reference_entries, key=_sort_key) if sort_references \
            else reference_entries
        for entry in entries:
            _reference_entry(doc, entry)

    from .workspace import refuse_cloud_save
    refuse_cloud_save(output_path, "the Word document")
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    doc.save(output_path)
    # Language tags of the work language, and file properties that carry the
    # title instead of the library's defaults (no name: anonymous marking).
    from .office_finalise import finalise
    title = re.search(r"^#\s+(.+)$", raw, flags=re.MULTILINE)
    # The Student ID belongs in the page header, where the marker is meant to
    # see it. It does not belong in the file properties, where it reads as a
    # deliberate signature. Author and last-modified-by are left empty.
    finalise(output_path, author="",
             title=title.group(1).strip() if title else "",
             subject=_real(module_code))

    if verbose:
        print("Word document saved: %s" % output_path)
        print("  Student ID in header: %s (no name)" % student_id)
        print("  Reference entries:    %d%s"
              % (len(reference_entries),
                 " (sorted alphabetically)" if sort_references else ""))
    return output_path
