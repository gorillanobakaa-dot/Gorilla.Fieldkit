"""The academic tools proven by making them work: real files built in a scratch folder and measured.

    fieldkit academic preflight            PREFLIGHT OK (exit 0) or PREFLIGHT NOT OK (exit 3), and which rule failed

Why (from the academic harness this generalises, preflight.py, 2026-10-10): the tools get edited by whatever AI the
student has to hand, and an AI asked to "improve" something quietly drops the parts it does not understand. Reading
the code cannot catch that; building a file and measuring it can. Every rule is a yes or no with a fixed answer,
nothing for a model to argue with. Run it after any change, and before trusting a copy someone else has touched.

Every rule follows the student's own settings (layout.current(): the country's page size, the work language), so
it holds for each of the seven countries, not for one student's.
"""
import os
import shutil
import tempfile
import zipfile

from . import layout

DRAFT = ("# A Case Analysis\n\n## Introduction\n\nBody text (Smith, 2020).\n\n## References\n\n"
         "Smith, J. (2020) *A book*. London: Sage. Available at: https://example.org/x\n")


def _page(path):
    import docx
    s = docx.Document(path).sections[0]
    return round(s.page_width.cm, 1), round(s.page_height.cm, 1)


def _draft(folder):
    p = os.path.join(folder, "draft.md")
    with open(p, "w", encoding="utf-8") as f:
        f.write(DRAFT)
    return p


class Rules:
    def __init__(self):
        self.rows = []

    def check(self, rule, ok, detail=""):
        self.rows.append({"rule": rule, "ok": bool(ok), "detail": "" if ok else str(detail)[:300]})
        return ok


def paper_and_layout(r, folder, lay):
    import docx
    from docx.oxml.ns import qn
    from .build_word import build_word_document
    from .build_templates import build_word_template
    from . import document_types as dt
    want = tuple(round(x, 1) for x in lay["page_cm"])
    built = os.path.join(folder, "page.docx")
    build_word_document(_draft(folder), built, "12345678", "M", verbose=False)
    r.check(f"a built document is {lay['page']}", _page(built) == want, _page(built))
    wrong = []
    for key in sorted(k for k in dt.DOCUMENT_TYPES if dt.DOCUMENT_TYPES[k].get("output", "docx") != "pptx"):
        p = os.path.join(folder, f"t_{key}.docx")
        try:
            build_word_template(key, p, verbose=False)
        except Exception as e:                                # a template that cannot be built is a failure
            wrong.append(f"{key}: {e}")
            continue
        if _page(p) != want:
            wrong.append(f"{key}: {_page(p)}")
    r.check(f"every Word template is {lay['page']}", not wrong, wrong)

    def sequence(path):
        out = []
        for p in docx.Document(path).paragraphs:
            if any(b.get(qn("w:type")) == "page" for b in p._p.findall(".//" + qn("w:br"))):
                out.append("<break>")
            if p.text.strip():
                out.append(p.style.name if p.style.name != "Normal" else p.text.strip()[:20])
        return out
    seq = sequence(built)
    r.check("the title is followed by a page break", seq[:2] == ["Title", "<break>"], seq[:4])
    r.check("the reference list starts a new page", "<break>" in seq[2:], seq)
    r.check("no blank page from two breaks in a row",
            not any(seq[i] == seq[i + 1] == "<break>" for i in range(len(seq) - 1)), seq)
    toc = os.path.join(folder, "toc.docx")
    build_word_document(_draft(folder), toc, "12345678", "M", verbose=False, contents_page=True)
    r.check("Word is told to fill in the contents page", b"updateFields" in zipfile.ZipFile(toc).read("word/settings.xml"))


def identity(r, folder):
    from .build_word import build_word_document, page_furniture
    from .office_finalise import finalise, read_properties
    built = os.path.join(folder, "id.docx")
    build_word_document(_draft(folder), built, "G1234567", "MOD101", verbose=False)
    props = read_properties(built)
    r.check("a built document has no author in its properties",
            not props.get("creator") and not props.get("lastModifiedBy"), props)
    r.check("...and no student number there either", "G1234567" not in " ".join(props.values()), props)
    header, numbered = page_furniture(built)
    r.check("the Student ID IS on the page header", "G1234567" in header, header)
    r.check("the pages are numbered", numbered)
    # Word writes the name back on every save: it must be caught and emptied
    stamped = os.path.join(folder, "stamped.docx")
    with zipfile.ZipFile(built) as src, zipfile.ZipFile(stamped, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                x = data.decode("utf-8")
                x = x.replace("<cp:lastModifiedBy></cp:lastModifiedBy>", "<cp:lastModifiedBy>A Student</cp:lastModifiedBy>")
                x = x.replace("<cp:lastModifiedBy/>", "<cp:lastModifiedBy>A Student</cp:lastModifiedBy>")
                data = x.encode("utf-8")
            dst.writestr(item, data)
    r.check("a name Word wrote back is seen", read_properties(stamped).get("lastModifiedBy") == "A Student",
            read_properties(stamped))
    finalise(stamped)
    r.check("...and emptied by finish", not read_properties(stamped).get("lastModifiedBy"), read_properties(stamped))


def templates(r, folder):
    import docx
    from .build_templates import build_word_template
    p = os.path.join(folder, "lean.docx")
    build_word_template("case-study", p, verbose=False)
    d = docx.Document(p)
    st = [x.style.name for x in d.paragraphs if x.text.strip()]
    r.check("no guidance paragraphs in a template", "Guidance" not in st, sorted(set(st)))
    r.check("the headings are there", "Heading 1" in st, sorted(set(st)))
    r.check("the template carries page numbers",
            any("PAGE" in x._p.xml for x in d.sections[0].footer.paragraphs))


def gates(r, lay):
    from .reference_auditor import audit_text
    from .dash_check import find_dashes
    from . import language
    faults, _ = find_dashes("Care is complex — and costly.\n")
    r.check("a long dash is found", len(faults) >= 1, faults)
    other = "ro" if lay["language"] != "ro" else "en"
    sample = {"ro": "Aceasta este o problema foarte importanta pentru asistentii medicali din spital.",
              "en": "This is a very important problem for the nurses who work in the hospital."}[other]
    res = language.check(sample + "\n", lay["language"], None)
    r.check("a sentence in another language fails the language check", not res["ok"], res["other"])
    if lay["language"] == "en":
        us = language.check("The organization met last week.\n", "en", "en-GB")
        r.check("the other spelling variant is named", bool(us["spelling"]), us["spelling"])
    rep = audit_text("Body (Okonkwo, 2019).\n\n## References\n\nSmith, J. (2020) *A*. L: S.\n")
    r.check("a citation with no entry is a failure", bool(rep["orphans"]) and rep["issues"] > 0, rep["orphans"])
    q = audit_text('Smith argues that "care should be tailored to the individual person rather than the condition" '
                   'in practice (Smith, 2020).\n\nHe also said "care should be designed around the whole person and '
                   'their family" here.\n\n## References\n\nSmith, J. (2020) *A*. L: S.\n')
    r.check("an uncited quotation warns and does not block", bool(q["quote_warnings"]) and not q["issues"],
            {"warn": q["quote_warnings"], "block": q["issues"]})


def language_tag(r, folder, lay):
    from .build_word import build_word_document
    from .office_finalise import languages
    built = os.path.join(folder, "lang.docx")
    build_word_document(_draft(folder), built, "12345678", "M", verbose=False)
    tags = languages(built)
    r.check(f"the document is tagged {lay['lang_tag']}", set(tags) == {lay["lang_tag"]}, tags)
    xml = b"".join(zipfile.ZipFile(built).read(n) for n in zipfile.ZipFile(built).namelist()
                   if n.endswith(".xml")).decode("utf-8", "replace")
    r.check("no long dash reaches the document", not any(d in xml for d in ("—", "–")))


def cloud(r):
    from .workspace import cloud_service
    wrong = [(p, cloud_service(p), want) for p, want in (
        (os.path.join("C:", os.sep, "Users", "x", "OneDrive", "Documents", "Exam_1"), "OneDrive"),
        (os.path.join("C:", os.sep, "Users", "x", "OneDrive - A University", "x"), "OneDrive"),
        (os.path.join("D:", os.sep, "Dropbox", "uni"), "Dropbox"),
        (os.path.join("C:", os.sep, "Users", "x", "Google Drive", "x"), "Google Drive")) if cloud_service(p) != want]
    r.check("a cloud folder is recognised wherever it sits in the path", not wrong, wrong)
    local = [p for p in (os.path.join("D:", os.sep, "uni", "Exam_1"),
                         os.path.join("C:", os.sep, "Users", "x", "Desktop", "onedriver_notes")) if cloud_service(p)]
    r.check("a local folder is not called cloud by mistake", not local, local)


def run(profile=None):
    lay = layout.current(profile)
    layout.apply(lay)
    r = Rules()
    folder = tempfile.mkdtemp(prefix="fk_preflight_")
    try:
        for name, step in (("page and layout", lambda: paper_and_layout(r, folder, lay)),
                           ("identity", lambda: identity(r, folder)), ("templates", lambda: templates(r, folder)),
                           ("checks that must block", lambda: gates(r, lay)),
                           ("language", lambda: language_tag(r, folder, lay)), ("cloud folders", lambda: cloud(r))):
            try:
                step()
            except Exception as e:
                r.check(f"{name} ran without crashing", False, f"{type(e).__name__}: {e}")
    finally:
        shutil.rmtree(folder, ignore_errors=True)
    failed = [x for x in r.rows if not x["ok"]]
    return {"ok": not failed, "rules": r.rows, "page": lay["page"], "language": lay["lang_tag"],
            "verdict": "PREFLIGHT OK" if not failed else "PREFLIGHT NOT OK",
            "next": "nothing to do" if not failed else "do not use the tools like this: something broke them; "
                                                        "git diff shows what changed"}


def lines(r):
    out = [f"  [{'OK  ' if x['ok'] else 'FAIL'}] {x['rule']}" + (f"\n         {x['detail']}" if x["detail"] else "")
           for x in r["rules"]]
    passed = sum(x["ok"] for x in r["rules"])
    return out + [f"{len(r['rules'])} checked, {passed} passed, {len(r['rules']) - passed} failed "
                  f"({r['page']}, {r['language']})", r["verdict"]]
